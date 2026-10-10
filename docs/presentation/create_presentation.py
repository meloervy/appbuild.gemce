import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
BG_COLOR = RGBColor(14, 16, 21)        # Deep Charcoal Navy
CARD_BG = RGBColor(24, 27, 36)         # Surface Card
CARD_BORDER = RGBColor(46, 53, 70)     # Card Border
PRIMARY_GREEN = RGBColor(61, 220, 132) # Gemce Mint Green
ACCENT_GREEN_BG = RGBColor(22, 45, 34) # Pill Background
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(165, 175, 190)
TEXT_ACCENT = RGBColor(141, 227, 181)
BLUE_ACCENT = RGBColor(96, 165, 250)
RED_ACCENT = RGBColor(248, 113, 113)

MOCKUPS_DIR = '/tmp/gemce_mockups_hi'
LOGO_PATH = '/home/mel-garcia/AndroidStudioProjects/Gemce/app/src/main/res/drawable/gemce_logo.png'

def set_slide_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_header(slide, badge_text, title_text, subtitle_text, header_width=Inches(11.7)):
    # Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(3.4), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT_GREEN_BG
    badge.line.color.rgb = PRIMARY_GREEN
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.word_wrap = False
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
    p_b = tf_b.paragraphs[0]
    p_b.text = badge_text
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = PRIMARY_GREEN
    p_b.font.name = "Arial"

    # Title & Subtitle Box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), header_width, Inches(1.15))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = title_text
    p1.font.size = Pt(25)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Arial"
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_MUTED
    p2.font.name = "Arial"
    p2.space_before = Pt(4)

def add_card(slide, left, top, width, height, title, items, border_color=CARD_BORDER, bg_color=CARD_BG, title_color=TEXT_WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    if title:
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.font.name = "Arial"
        p_title.space_after = Pt(6)
    
    first = True if not title else False
    for item in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Arial"
        p.space_after = Pt(3)
    return card

# ==========================================
# SLIDE 1: COVER SLIDE
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

if os.path.exists(LOGO_PATH):
    slide1.shapes.add_picture(LOGO_PATH, Inches(0.8), Inches(1.0), Inches(1.2), Inches(1.2))

tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(8.2), Inches(3.6))
tf1 = tb1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

p1 = tf1.paragraphs[0]
p1.text = "Gemce"
p1.font.size = Pt(56)
p1.font.bold = True
p1.font.color.rgb = TEXT_WHITE
p1.font.name = "Arial"

p2 = tf1.add_paragraph()
p2.text = "AI study buddy that works anywhere you go."
p2.font.size = Pt(22)
p2.font.bold = True
p2.font.color.rgb = PRIMARY_GREEN
p2.font.name = "Arial"
p2.space_before = Pt(8)

p3 = tf1.add_paragraph()
p3.text = "The 100% offline, privacy-first Android study platform powered by on-device Google Gemma 4 E2B via LiteRT-LM. Eliminating the connectivity divide in education."
p3.font.size = Pt(14)
p3.font.color.rgb = TEXT_MUTED
p3.font.name = "Arial"
p3.space_before = Pt(14)

badges = ["🔒 100% Offline", "⚡ Gemma 4 E2B", "📚 Verbatim Citations", "🧠 SuperMemo SM-2"]
bx = Inches(0.8)
by = Inches(5.3)
for b_text in badges:
    b_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(2.0), Inches(0.42))
    b_shape.fill.solid()
    b_shape.fill.fore_color.rgb = RGBColor(26, 31, 42)
    b_shape.line.color.rgb = RGBColor(50, 60, 80)
    b_shape.line.width = Pt(1)
    tf_b = b_shape.text_frame
    tf_b.word_wrap = False
    p = tf_b.paragraphs[0]
    p.text = b_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(220, 230, 245)
    p.font.name = "Arial"
    bx += Inches(2.12)

mockup1 = f'{MOCKUPS_DIR}/01_start.png'
if os.path.exists(mockup1):
    slide1.shapes.add_picture(mockup1, Inches(9.8), Inches(0.8), Inches(2.86), Inches(5.9))

# ==========================================
# SLIDE 2: THE PROBLEM & SOLUTION
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_header(slide2, "01 · PROBLEM & SOLUTION", "Why Cloud AI Fails — And How Gemce Solves It", "Eliminating connectivity barriers, recurring paywalls, and privacy risks.")

p_items = [
    "• Fails in Airplane Mode: Cloud AI drops immediately in transit, dorms, or rural areas.",
    "• Cost Barriers: $20/month subscriptions lock out everyday students.",
    "• Privacy Concerns: Uploads private notes and school materials to cloud servers."
]
add_card(slide2, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.1), "❌ The Problem: Cloud AI Limitations", p_items, border_color=RGBColor(80, 45, 50), bg_color=RGBColor(28, 20, 24), title_color=RED_ACCENT)

s_items = [
    "• 100% Local Inference: Runs Google Gemma 4 on-device with zero data connection.",
    "• Free & Sovereign: No subscriptions, no telemetry, and complete data privacy.",
    "• Verifiable Proof: Every card and question cites exact sentences from student notes."
]
add_card(slide2, Inches(6.9), Inches(2.2), Inches(5.6), Inches(4.1), "✅ The Solution: Gemce On-Device AI", s_items, border_color=RGBColor(35, 75, 55), bg_color=RGBColor(18, 30, 26), title_color=PRIMARY_GREEN)

banner2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.55))
banner2.fill.solid()
banner2.fill.fore_color.rgb = ACCENT_GREEN_BG
banner2.line.color.rgb = PRIMARY_GREEN
banner2.line.width = Pt(1)
p_ban = banner2.text_frame.paragraphs[0]
p_ban.text = "💡 Fundamental Principle: Quality academic assistance should never depend on mobile data or costly subscriptions."
p_ban.alignment = PP_ALIGN.CENTER
p_ban.font.size = Pt(12)
p_ban.font.bold = True
p_ban.font.color.rgb = TEXT_ACCENT

# ==========================================
# SLIDE 3: ADAPTIVE STUDY HUB
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_header(slide3, "02 · ADAPTIVE STUDY HUB", "Personalized Study Paths Tailored to Your Goals", "Diagnoses student intent and calibrates academic rigor, grade levels, and language.", header_width=Inches(8.2))

hub_cards = [
    ("🎯 Goal-Oriented Learning Modes", [
        "• Diagnoses study intent: Rapid Recall, Exam Practice, Deep Problem Solving, or Reviewers.",
        "• Automatically prescribes optimal study formats based on topic familiarity."
    ]),
    ("🎓 Academic Level & Major Calibration", [
        "• Calibrates technical depth from Junior/Senior High tracks to specialized College majors.",
        "• Tailors difficulty for Accountancy, Nursing, Engineering, CS, and Humanities."
    ]),
    ("🌐 Multilingual Academic Mentoring", [
        "• Speaks English, Filipino (Tagalog), Taglish, and Cebuano (Bisaya).",
        "• Explains tough concepts in the student's natural, conversational tongue."
    ])
]

cy = Inches(2.2)
for title, items in hub_cards:
    add_card(slide3, Inches(0.8), cy, Inches(7.4), Inches(1.35), title, items)
    cy += Inches(1.5)

mockup3 = f'{MOCKUPS_DIR}/02_study_hub.png'
if os.path.exists(mockup3):
    slide3.shapes.add_picture(mockup3, Inches(9.2), Inches(1.2), Inches(2.85), Inches(5.85))

# ==========================================
# SLIDE 4: SPACED REPETITION FLASHCARDS
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_header(slide4, "03 · ACTIVE RECALL", "Evidence-Grounded Flashcards with SuperMemo SM-2", "Spaced repetition paired with zero-hallucination source grounding.", header_width=Inches(6.0))

flash_cards = [
    ("🧠 SuperMemo SM-2 Spaced Repetition", [
        "• Dynamically schedules review intervals (Again, Hard, Good, Easy).",
        "• Optimizes long-term memory retention before examinations."
    ]),
    ("🚫 Strict Anti-Hallucination Prompts", [
        "• Forbids generic meta-questions like 'What did the user intend to study?'.",
        "• Tests concrete academic mechanisms: Mitochondria, ATP synthesis, rock types."
    ]),
    ("📖 Verbatim Source Provenance", [
        "• Highlights verbatim source sentences on the card back.",
        "• Guarantees 100% factual accuracy with zero AI hallucinations."
    ])
]

cy = Inches(2.2)
for title, items in flash_cards:
    add_card(slide4, Inches(0.8), cy, Inches(5.6), Inches(1.35), title, items)
    cy += Inches(1.5)

mockup4_q = f'{MOCKUPS_DIR}/04_card_front.png'
mockup4_a = f'{MOCKUPS_DIR}/05_card_back.png'
if os.path.exists(mockup4_q):
    slide4.shapes.add_picture(mockup4_q, Inches(7.0), Inches(1.2), Inches(2.8), Inches(5.8))
if os.path.exists(mockup4_a):
    slide4.shapes.add_picture(mockup4_a, Inches(10.1), Inches(1.2), Inches(2.8), Inches(5.8))

# ==========================================
# SLIDE 5: INTERACTIVE PRACTICE QUIZZES
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_header(slide5, "04 · EXAM PREPARATION", "Interactive Practice Quizzes & Instant Evaluation", "4-option MCQs with instant scoring and explanatory rationales.", header_width=Inches(6.0))

quiz_cards = [
    ("🎯 4-Option Multiple Choice Assessments", [
        "• Rigorous MCQs testing conceptual mastery and analytical discrimination.",
        "• Challenges students to apply principles rather than passively re-read."
    ]),
    ("⚡ Instant Grading & Rationale Explanations", [
        "• Instantly highlights incorrect selections in red and correct choices in teal.",
        "• Explanatory rationale box details the exact conceptual reason."
    ]),
    ("📚 100% Material Grounding", [
        "• Every question, distractor, and rationale derives directly from student notes.",
        "• Eliminates invalid test keys and erroneous study facts."
    ])
]

cy = Inches(2.2)
for title, items in quiz_cards:
    add_card(slide5, Inches(0.8), cy, Inches(5.6), Inches(1.35), title, items)
    cy += Inches(1.5)

mockup5_live = f'{MOCKUPS_DIR}/06_quiz_live.png'
mockup5_checked = f'{MOCKUPS_DIR}/07_quiz_checked.png'
if os.path.exists(mockup5_live):
    slide5.shapes.add_picture(mockup5_live, Inches(7.0), Inches(1.2), Inches(2.8), Inches(5.8))
if os.path.exists(mockup5_checked):
    slide5.shapes.add_picture(mockup5_checked, Inches(10.1), Inches(1.2), Inches(2.8), Inches(5.8))

# ==========================================
# SLIDE 6: ASK GEMCE ACADEMIC TUTOR
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_header(slide6, "05 · INTERACTIVE TUTORING", "Real-Time Academic Mentor in Airplane Mode", "On-device conversational AI explaining complex mechanisms on demand.", header_width=Inches(8.2))

chat_cards = [
    ("🌍 Universal Multidisciplinary Grounding", [
        "• Mentors across STEM, Calculus, Discrete Math, History, Rizal, Law, and Medicine.",
        "• Unpacks tough mechanisms with step-by-step conceptual walkthroughs."
    ]),
    ("📑 Rich Markdown & Code Formatting", [
        "• Native Jetpack Compose formatting with tables, bulleted lists, and formulas.",
        "• Clean visual hierarchy avoids overwhelming walls of plain text."
    ]),
    ("💾 Multi-Session Local Memory & Tabs", [
        "• Conversation history persisted securely on device with quick topic tabs.",
        "• Complete data privacy: student questions never leave device memory."
    ])
]

cy = Inches(2.2)
for title, items in chat_cards:
    add_card(slide6, Inches(0.8), cy, Inches(7.4), Inches(1.35), title, items)
    cy += Inches(1.5)

mockup6 = f'{MOCKUPS_DIR}/08_chat.png'
if os.path.exists(mockup6):
    slide6.shapes.add_picture(mockup6, Inches(9.2), Inches(1.2), Inches(2.85), Inches(5.85))

# ==========================================
# SLIDE 7: TECHNICAL ARCHITECTURE & ROADMAP
# ==========================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_header(slide7, "06 · ARCHITECTURE & ROADMAP", "Engineered for Performance, Ready to Deploy", "Native Android architecture powered by Google LiteRT-LM and open-source delivery.", header_width=Inches(8.2))

tech_cards = [
    ("🤖 Google LiteRT-LM & Gemma 4 E2B", [
        "• Quantized 2.4 GB weights executed on-device with OpenCL / Adreno GPU acceleration.",
        "• Sub-second latency with smooth streaming generation entirely offline."
    ]),
    ("🎨 Native Jetpack Compose & Offline Room DB", [
        "• 100% modern Kotlin with Material 3 and packaged typography (Calibri, Serif, Mono).",
        "• Offline OCR via ML Kit v2 and pure local PDF ingestion."
    ]),
    ("📲 Shipped MVP & Public GitHub Distribution", [
        "• Verified on Android 14 hardware (OPPO Reno12 F 5G).",
        "• Public Repository: github.com/meloervy/appbuild.gemce with installable APK."
    ])
]

cy = Inches(2.2)
for title, items in tech_cards:
    add_card(slide7, Inches(0.8), cy, Inches(7.4), Inches(1.35), title, items)
    cy += Inches(1.5)

mockup7 = f'{MOCKUPS_DIR}/09_about.png'
if os.path.exists(mockup7):
    slide7.shapes.add_picture(mockup7, Inches(9.2), Inches(1.2), Inches(2.85), Inches(5.85))

# Save presentation
out_pptx = '/home/mel-garcia/AndroidStudioProjects/Gemce/docs/presentation/Gemce_Pitch_Presentation.pptx'
prs.save(out_pptx)
prs.save('/home/mel-garcia/AndroidStudioProjects/Gemce/Gemce_Pitch_Presentation.pptx')
print(f'7-slide presentation saved to {out_pptx}')
